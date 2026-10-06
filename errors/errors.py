"""
parser.parser_engine
====================

ParserEngine is the fifth pipeline stage. It reads the repository
inventory published by Module 4's RepositoryManager, runs every selected
notebook through the full parser stack (Notebook Reader -> Magic Parser ->
Regex Parser -> AST Parser -> SQL Parser -> Dependency Extractor), builds
the repo-wide dependency graph, and generates the final
MigrationKnowledgeModel.

The migration.notebook_list is authoritative when provided.
"""

from __future__ import annotations

import json
from pathlib import Path
from typing import Dict, List

from common.exceptions import ParserError
from common.logging_config import get_logger
from orchestrator.context import PipelineContext
from orchestrator.interfaces import PipelineStage
from parser.ast_parser import ASTParser
from parser.dependency_extractor import DependencyExtractor, NotebookDependency
from parser.dependency_graph import DependencyGraphBuilder
from parser.knowledge_model import KnowledgeModelGenerator, NotebookParseResult
from parser.magic_parser import MagicCommandParser
from parser.migration_analyzer import MigrationAnalyzer
from parser.notebook_reader import NotebookReader
from parser.regex_parser import RegexParser
from parser.sql_parser import SQLParser
from repository.inventory import NotebookLanguage, NotebookFile


logger = get_logger(__name__)

_PROJECT_ROOT = Path(__file__).resolve().parents[2]


class ParserEngine(PipelineStage):
    """
    Parses selected notebooks from the synced repository
    (from context.artifacts["repository"]) and produces a
    MigrationKnowledgeModel.
    """

    name = "ParserEngine"

    def __init__(self):
        self._reader = NotebookReader()
        self._magic_parser = MagicCommandParser()
        self._regex_parser = RegexParser()
        self._ast_parser = ASTParser()
        self._sql_parser = SQLParser()
        self._knowledge_model_generator = KnowledgeModelGenerator()

    def run(self, context: PipelineContext) -> PipelineContext:
        repository = context.get_artifact("repository")

        if not repository:
            raise ParserError(
                "ParserEngine requires context.artifacts['repository'] to be "
                "populated (run RepositoryManager first)"
            )

        if not context.config:
            raise ParserError(
                "ParserEngine requires context.config to be populated"
            )

        # ------------------------------------------------------------
        # Repository inventory
        # ------------------------------------------------------------
        all_repository_entries = repository.notebooks

        known_paths = list(
            repository.known_notebook_paths
            or [
                nb.relative_path
                for nb in all_repository_entries
            ]
        )

        dependency_extractor = DependencyExtractor(
            known_notebook_paths=known_paths
        )

        # ------------------------------------------------------------
        # Migration scope
        # ------------------------------------------------------------
        selected_paths = self._load_migration_scope(
            context,
            known_paths,
        )

        normalized_selected = {
            self._normalize_scope_path(path, context)
            for path in selected_paths
        }

        # ------------------------------------------------------------
        # Match requested notebooks against repository inventory.
        #
        # The normalization allows all of these forms to match:
        #
        # /PFL/Delta-Lake/foo/notebook
        #
        # /PFL/Delta-Lake/foo/notebook.py
        #
        # /Workspace/Users/user/repo/PFL/Delta-Lake/foo/notebook
        #
        # /Workspace/Users/user/repo/PFL/Delta-Lake/foo/notebook.py
        # ------------------------------------------------------------
        notebook_entries = [
            nb
            for nb in all_repository_entries
            if self._normalize_scope_path(
                nb.relative_path,
                context,
            ) in normalized_selected
        ]

        found_normalized = {
            self._normalize_scope_path(
                nb.relative_path,
                context,
            )
            for nb in notebook_entries
        }

        missing_requested = [
            requested
            for requested in selected_paths
            if self._normalize_scope_path(
                requested,
                context,
            ) not in found_normalized
        ]

        # ------------------------------------------------------------
        # Publish migration scope information.
        # ------------------------------------------------------------
        context.set_artifact(
            "migration_scope",
            {
                "requested_notebooks": selected_paths,
                "found_notebooks": [
                    nb.relative_path
                    for nb in notebook_entries
                ],
                "missing_notebooks": missing_requested,
                "all_source_notebooks": known_paths,
            },
        )

        if missing_requested:
            logger.warning(
                "%d requested notebook(s) were not found in "
                "source repository: %s",
                len(missing_requested),
                missing_requested,
            )

        # ------------------------------------------------------------
        # IMPORTANT:
        # If an explicit migration list was supplied but absolutely
        # nothing matched, stop the pipeline.
        #
        # This prevents the previous false-success scenario:
        #
        #   5 requested
        #   0 parsed
        #   0 transformed
        #   DEPLOYED
        # ------------------------------------------------------------
        if selected_paths and not notebook_entries:
            raise ParserError(
                "Migration scope contains requested notebooks, but "
                "ParserEngine could not match any requested notebook "
                "againt the repositoryManager inventory. "
                "Check workspace-root/path normalization."
            )

        # ------------------------------------------------------------
        # Migration configuration
        # ------------------------------------------------------------
        migration_cfg = context.config.get("migration", {}) or {}

        mapping = dict(
            migration_cfg.get(
                "container_bucket_mapping",
                {},
            )
            or {}
        )

        mapping_file = migration_cfg.get(
            "container_bucket_mapping_file"
        )

        # ------------------------------------------------------------
        # Optional external container -> bucket mapping
        # ------------------------------------------------------------
        if mapping_file:
            mapping_path = Path(mapping_file)

            if not mapping_path.is_absolute():
                mapping_path = _PROJECT_ROOT / mapping_path

            if not mapping_path.exists():
                raise ParserError(
                    f"Container-to-bucket mapping file does not exist: "
                    f"{mapping_path}"
                )

            try:
                # --------------------------------------------
                # YAML
                # --------------------------------------------
                if mapping_path.suffix.lower() in {
                    ".yaml",
                    ".yml",
                }:
                    import yaml

                    loaded = yaml.safe_load(
                        mapping_path.read_text(
                            encoding="utf-8"
                        )
                    ) or {}

                    if not isinstance(loaded, dict):
                        raise ParserError(
                            "Container-to-bucket YAML mapping "
                            "must be a mapping"
                        )

                    mapping.update(loaded)

                # --------------------------------------------
                # CSV / TXT
                # --------------------------------------------
                elif mapping_path.suffix.lower() in {
                    ".csv",
                    ".txt",
                }:
                    import csv

                    with mapping_path.open(
                        "r",
                        encoding="utf-8-sig",
                        newline="",
                    ) as fh:
                        for row in csv.reader(fh):
                            if (
                                len(row) >= 2
                                and row[0].strip().lower()
                                not in {
                                    "container",
                                    "azure container",
                                }
                            ):
                                mapping[
                                    row[0].strip()
                                ] = row[1].strip()

                # --------------------------------------------
                # Excel
                # --------------------------------------------
                elif mapping_path.suffix.lower() in {
                    ".xlsx",
                    ".xlsm",
                }:
                    from openpyxl import load_workbook

                    wb = load_workbook(
                        mapping_path,
                        read_only=True,
                        data_only=True,
                    )

                    rows = list(
                        wb.active.iter_rows(
                            values_only=True
                        )
                    )

                    for row in rows[1:]:
                        if (
                            len(row) >= 2
                            and row[0]
                            and row[1]
                        ):
                            mapping[
                                str(row[0]).strip()
                            ] = str(row[1]).strip()

                else:
                    raise ParserError(
                        "Unsupported mapping file format: "
                        f"{mapping_path.suffix}"
                    )

            except ImportError as exc:
                raise ParserError(
                    "Required dependency is missing to read "
                    f"mapping file {mapping_path}"
                ) from exc

        # ------------------------------------------------------------
        # Migration analyzer
        # ------------------------------------------------------------
        analyzer = MigrationAnalyzer(
            container_bucket_mapping=mapping,
            workspace_root=migration_cfg.get(
                "workspace_root",
                "${WORKSPACE_ROOT}",
            ),
            referred_notebook_path_root=migration_cfg.get(
                "referred_notebook_path_root",
                "${TARGET_WORKSPACE_ROOT}",
            ),
            source_repo_path=(
                context.config.get("source", {})
                or {}
            ).get(
                "repo_path",
                "",
            ),
            bucket_variable=migration_cfg.get(
                "bucket_variable",
                "bucket",
            ),
        )

        # ------------------------------------------------------------
        # Parse notebooks
        # ------------------------------------------------------------
        notebook_results: List[NotebookParseResult] = []
        all_dependencies: List[NotebookDependency] = []
        migration_analyses = []

        for entry in notebook_entries:
            result, deps = self._parse_one_notebook(
                entry,
                dependency_extractor,
                analyzer,
            )

            notebook_results.append(result)
            all_dependencies.extend(deps)

            migration_analyses.append(
                result.migration_analysis.to_dict()
                if result.migration_analysis
                else {
                    "notebook": result.relative_path,
                    "edits": [],
                }
            )

        # ------------------------------------------------------------
        # Safety check:
        # explicit migration list + zero parsed notebooks
        # ------------------------------------------------------------
        if selected_paths and not notebook_results:
            raise ParserError(
                "Migration scope contains requested notebooks, but "
                "ParserEngine parsed zero notebooks. "
            )

        # ------------------------------------------------------------
        # Determine whether dependencies themselves are part of the
        # explicit migration list.
        # ------------------------------------------------------------
        selected_set = {
            self._normalize_scope_path(
                path,
                context,
            )
            for path in selected_paths
        }

        for dependency in all_dependencies:
            dependency.target_in_migration_list = (
                self._normalize_scope_path(
                    dependency.resolved_target,
                    context,
                )
                in selected_set
                if dependency.resolved_target
                else False
            )

        # ------------------------------------------------------------
        # Dependency graph
        # ------------------------------------------------------------
        graph_builder = DependencyGraphBuilder()

        graph = graph_builder.build(
            known_paths,
            all_dependencies,
        )

        graph_summary = graph_builder.summarize(
            graph,
            all_dependencies,
        )

        # ------------------------------------------------------------
        # Knowledge model
        # ------------------------------------------------------------
        model = self._knowledge_model_generator.generate(
            repo_name=repository.repo_name,
            notebook_results=notebook_results,
            dependencies=all_dependencies,
            graph_summary=graph_summary,
        )

        model.migration_scope = context.get_artifact(
            "migration_scope",
            {},
        )

        # ------------------------------------------------------------
        # Write knowledge model
        # ------------------------------------------------------------
        output_path = self._write_knowledge_model(
            context,
            model.to_dict(),
        )

        # ------------------------------------------------------------
        # Publish artifacts
        # ------------------------------------------------------------
        context.set_artifact(
            "knowledge_model",
            model.to_dict(),
        )

        context.set_artifact(
            "dependency_graph_summary",
            graph_summary.to_dict(),
        )

        context.set_artifact(
            "dependency_graph_dot",
            graph_builder.to_dot(graph),
        )

        context.set_artifact(
            "migration_analyses",
            migration_analyses,
        )

        context.set_artifact(
            "parsed_notebook_count",
            len(notebook_results),
        )

        context.metadata["parser_complete"] = True
        context.metadata["knowledge_model_path"] = str(
            output_path
        )

        logger.info(
            "Parsed %d notebook(s); %d dependency edge(s), "
            "%d azure construct(s), %d manual review item(s). "
            "Knowledge model written to %s",
            len(notebook_results),
            len(model.dependencies["edges"]),
            len(model.azure_constructs),
            len(model.manual_review_items),
            output_path,
        )

        return context

    # ------------------------------------------------------------------
    # Migration scope
    # ------------------------------------------------------------------

    @staticmethod
    def _load_migration_scope(
        context: PipelineContext,
        known_paths: List[str],
    ) -> List[str]:

        migration_cfg = (
            context.config.get("migration", {})
            or {}
        )

        # ------------------------------------------------------------
        # Preferred input: explicit inline list.
        #
        # The list is authoritative.
        # Notebooks discovered as dependencies are NOT automatically
        # added to the migration scope.
        # ------------------------------------------------------------
        inline_list = migration_cfg.get(
            "notebook_list"
        ) or []

        if inline_list:
            if not isinstance(
                inline_list,
                (list, tuple),
            ):
                raise ParserError(
                    "migration.notebook_list must be a list"
                )

            return list(
                dict.fromkeys(
                    str(v).strip()
                    for v in inline_list
                    if str(v).strip()
                )
            )

        # ------------------------------------------------------------
        # Backward-compatible file input.
        # ------------------------------------------------------------
        list_file = migration_cfg.get(
            "notebook_list_file"
        )

        if not list_file:
            return list(known_paths)

        path = Path(list_file)

        if not path.is_absolute():
            path = _PROJECT_ROOT / path

        if not path.exists():
            raise ParserError(
                "Migration notebook list file does not exist: "
                f"{path}"
            )

        suffix = path.suffix.lower()
        values = []

        # ------------------------------------------------------------
        # CSV / TXT
        # ------------------------------------------------------------
        if suffix in {
            ".csv",
            ".txt",
        }:
            import csv

            with path.open(
                "r",
                encoding="utf-8-sig",
                newline="",
            ) as fh:

                rows = csv.reader(fh)

                column = str(
                    migration_cfg.get(
                        "notebook_path_column",
                        "Notebook Path",
                    )
                )

                for row in rows:
                    if not row:
                        continue

                    if not row[0].strip():
                        continue

                    value = row[0].strip()

                    if value == column:
                        continue

                    values.append(value)

        # ------------------------------------------------------------
        # XLSX / XLSM
        # ------------------------------------------------------------
        elif suffix in {
            ".xlsx",
            ".xlsm",
        }:
            try:
                from openpyxl import load_workbook
            except ImportError as exc:
                raise ParserError(
                    "openpyxl is required to read an Excel "
                    "migration notebook list"
                ) from exc

            wb = load_workbook(
                path,
                read_only=True,
                data_only=True,
            )

            ws = wb.active

            rows = list(
                ws.iter_rows(
                    values_only=True
                )
            )

            column = str(
                migration_cfg.get(
                    "notebook_path_column",
                    "Notebook Path",
                )
            )

            header = [
                str(v).strip()
                if v is not None
                else ""
                for v in (
                    rows[0]
                    if rows
                    else []
                )
            ]

            try:
                idx = header.index(column)
            except ValueError:
                idx = 0

            values = [
                str(row[idx]).strip()
                for row in rows[1:]
                if (
                    len(row) > idx
                    and row[idx]
                )
            ]

        else:
            raise ParserError(
                "Unsupported migration notebook list "
                f"format: {suffix}. "
                "Use CSV, TXT, XLSX or XLSM."
            )

        return list(
            dict.fromkeys(values)
        )

    # ------------------------------------------------------------------
    # Path normalization
    # ------------------------------------------------------------------

    @staticmethod
    def _normalize_scope_path(
        path: str,
        context: PipelineContext,
    ) -> str:
        """
        Normalize notebook paths so the following forms are treated
        as the same notebook:

            /PFL/Delta-Lake/Common/utils

            /PFL/Delta-Lake/Common/utils.py

            /Workspace/Users/user/repo/PFL/Delta-Lake/Common/utils

            /Workspace/Users/user/repo/PFL/Delta-Lake/Common/utils.py

        The returned value is:

            repository-relative
            lowercase
            extensionless
        """

        value = str(
            path or ""
        ).strip().replace(
            "\\",
            "/",
        )

        source = (
            context.config.get("source", {})
            if context.config
            else {}
        )

        # Prefer workspace_path, then repo_path.
        root = str(
            source.get("workspace_path")
            or source.get("repo_path")
            or "/"
        ).strip().replace(
            "\\",
            "/",
        )

        root = root.rstrip("/")

        # Always work with a leading slash internally.
        value = "/" + value.strip("/")

        root_normalized = (
            "/" + root.strip("/")
            if root.strip("/")
            else ""
        )

        # ------------------------------------------------------------
        # Remove configured workspace root.
        #
        # Example:
        #
        # /Workspace/Users/user/repo/PFL/foo
        #
        # becomes:
        #
        # /PFL/foo
        # ------------------------------------------------------------
        if root_normalized:
            root_lower = root_normalized.lower()
            value_lower = value.lower()

            if value_lower == root_lower:
                value = "/"

            elif value_lower.startswith(
                root_lower + "/"
            ):
                value = value[
                    len(root_normalized):
                ]

        # ------------------------------------------------------------
        # Normalize again after root removal.
        # ------------------------------------------------------------
        value = "/" + value.strip("/")

        # ------------------------------------------------------------
        # Remove local source extensions.
        #
        # Databricks workspace paths are normally extensionless,
        # while exported source files contain .py/.sql/.scala.
        # ------------------------------------------------------------
        lower_value = value.lower()

        for extension in (
            ".py",
            ".sql",
            ".scala",
        ):
            if lower_value.endswith(extension):
                value = value[
                    :-len(extension)
                ]
                break

        # ------------------------------------------------------------
        # Final canonical representation.
        # ------------------------------------------------------------
        return value.strip("/").lower()

    # ------------------------------------------------------------------
    # Notebook parser
    # ------------------------------------------------------------------

    def _parse_one_notebook(
        self,
        entry: NotebookFile,
        dependency_extractor: DependencyExtractor,
        analyzer: MigrationAnalyzer,
    ):
        relative_path = entry.relative_path
        absolute_path = Path(
            entry.absolute_path
        )
        category = entry.category

        # ------------------------------------------------------------
        # Read notebook
        # ------------------------------------------------------------
        try:
            notebook_source = self._reader.read(
                relative_path,
                str(absolute_path),
            )

        except ParserError as exc:
            logger.warning(
                "Notebook '%s' could not be read: %s",
                relative_path,
                exc,
            )

            result = NotebookParseResult(
                relative_path=relative_path,
                category=category,
                language=entry.language.unknown,
                cell_count=0,
                parse_error=str(exc),
            )

            return result, []

        # ------------------------------------------------------------
        # Magic commands
        # ------------------------------------------------------------
        magic_commands = self._magic_parser.parse(
            notebook_source
        )

        # ------------------------------------------------------------
        # Azure / cloud construct detection
        # ------------------------------------------------------------
        azure_constructs = self._regex_parser.parse(
            notebook_source.full_text
        )

        # ------------------------------------------------------------
        # Migration analysis
        # ------------------------------------------------------------
        migration_analysis = analyzer.analyze(
            relative_path,
            notebook_source.full_text,
        )

        # ------------------------------------------------------------
        # AST parsing for Python
        # ------------------------------------------------------------
        ast_findings = None
        parse_error = None

        if (
            notebook_source.language
            == NotebookLanguage.PYTHON
        ):
            try:
                ast_findings = self._ast_parser.parse(
                    notebook_source.full_text,
                    relative_path,
                )

            except ParserError as exc:
                parse_error = str(exc)

                logger.warning(
                    "AST parse failed for '%s': %s",
                    relative_path,
                    exc,
                )

        # ------------------------------------------------------------
        # SQL parsing
        # ------------------------------------------------------------
        sql_findings = self._extract_sql_findings(
            notebook_source,
            magic_commands,
        )

        # ------------------------------------------------------------
        # Build parse result
        # ------------------------------------------------------------
        result = NotebookParseResult(
            relative_path=relative_path,
            category=category,
            language=notebook_source.language.value,
            cell_count=notebook_source.cell_count,
            magic_commands=magic_commands,
            azure_constructs=azure_constructs,
            ast_findings=ast_findings,
            sql_findings=sql_findings,
            parse_error=parse_error,
            migration_analysis=migration_analysis,
        )

        # ------------------------------------------------------------
        # Dependency extraction
        # ------------------------------------------------------------
        dependencies = dependency_extractor.extract(
            relative_path,
            magic_commands,
            ast_findings,
        )

        return result, dependencies

    # ------------------------------------------------------------------
    # SQL findings
    # ------------------------------------------------------------------

    def _extract_sql_findings(
        self,
        notebook_source,
        magic_commands,
    ):
        sql_text_parts: List[str] = []

        if (
            notebook_source.language
            == NotebookLanguage.SQL
        ):
            sql_text_parts.append(
                notebook_source.full_text
            )

        sql_text_parts.extend(
            m.argument
            for m in magic_commands
            if m.magic_type == "sql"
        )

        if not sql_text_parts:
            return None

        return self._sql_parser.parse(
            "\n".join(sql_text_parts)
        )

    # ------------------------------------------------------------------
    # Knowledge model output
    # ------------------------------------------------------------------

    def _write_knowledge_model(
        self,
        context: PipelineContext,
        model_dict: Dict,
    ) -> Path:

        output_dir_config = context.config[
            "output"
        ][
            "knowledge_model_dir"
        ]

        output_dir = Path(
            output_dir_config
        )

        if not output_dir.is_absolute():
            output_dir = (
                _PROJECT_ROOT
                / output_dir
            )

        output_dir.mkdir(
            parents=True,
            exist_ok=True,
        )

        output_path = (
            output_dir
            / "MigrationKnowledgeModel.json"
        )

        output_path.write_text(
            json.dumps(
                model_dict,
                indent=2,
                default=str,
            ),
            encoding="utf-8",
        )

        return output_path
